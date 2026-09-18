#!/usr/bin/env python3
"""
gerar_relper.py — a fonte RELPER: o formulario periodico decretado pelo EM da brigada.

Duas remessas por dia (MATUTINO ~06h, VESPERTINO ~18h) por OM, uma linha por
subunidade. Cinco OM devolvem o .xlsx; duas (GAC51 e CAV51) imprimem,
preenchem, assinam e digitalizam — mesmo formulario, modalidade PDF, e os
numeros passam por OCR antes de chegar aos mesmos campos.

O que varia de proposito, porque varia no campo:
  - a celula 'Fracao' e digitada a mao: '1a Cia', '1ª Cia Fuz', '1 Cia Fuz/511'
  - a coluna 'Nec Prio' e texto livre com abreviatura: 'Fz', 'Mun 7,62', 'Cl III'
  - duas OM usam a VERSAO 2 do formulario ('Comb %' no lugar de 'Comb Pctl',
    mais a coluna 'Agua L'); uma OM acrescentou 'Obs' por conta propria —
    coluna que o formulario nao previa vai para ATRIBUTO_EXTRA_TXT (sobra)

Cada arquivo sai com o sidecar .json: operacao, OM remetente, data-hora, turno.

Defeitos plantados: turno_ausente (a OM nao remeteu), abreviatura_trocada
('Fuz' no lugar de 'Fz': material vira pessoal), pasta_errada (um arquivo
pousou em landing/perseu_2023/, e o sidecar diz PERSEU_2024).
"""

from datetime import timedelta

from openpyxl import Workbook

from comum import Contexto, assinar, desenhar_tabela, escanear, escrever_sidecar, fonte, iso, local, pagina, salvar_pdf, texto_centrado

CABECALHO_V1 = ["Fração", "Ef Prev", "Ef Pres", "Baixas Cmb", "Baixas N Cmb", "Evacuados", "Vtr Op", "Vtr Total",
                "Comb Pctl", "Sup Cl I", "Sup Cl III", "Sup Cl V", "Sup Cl VIII", "Nec Prio"]
CABECALHO_V2 = [c if c != "Comb Pctl" else "Comb %" for c in CABECALHO_V1] + ["Agua L"]
OM_PDF = {"GAC51", "CAV51"}
OM_V2 = {"BI513", "GAC51"}
OM_COLUNA_EXTRA = {"BI512": "Obs"}
NECESSIDADES = ["Fz", "Mun 7,62", "Cl III", "Agua", "Cl I", "Racao", "Vtr", "Sup Cl VIII", "Pneu", "Bateria",
                "Nada a relatar", "Nada a relatar", "Nada a relatar"]
OBS = ["", "", "Vtr em Mnt", "Aguardando Cl V", "Ef em instrucao", "Sem alteracao"]
TURNOS = [("MATUTINO", 6), ("VESPERTINO", 18)]


def _linhas_da_om(ctx: Contexto, om_cod: str) -> list[dict]:
    """Quem preenche uma linha: as subunidades; se a OM nao tem SU, as fracoes diretas."""
    su = [u for u in ctx.filhas(om_cod) if u["ESCALAO_COD"] == "SU"]
    return su or ctx.filhas(om_cod)


def _grafia_fracao(ctx: Contexto, u: dict) -> str:
    """A celula digitada a mao: varias grafias para a mesma unidade."""
    rng = ctx.rng
    sigla = u["UNIDADE_SIGLA"]
    cabeca = sigla.split("/")[0]                      # '1a Cia Fuz'
    om_num = ctx.unidade(ctx.om_de(u["UNIDADE_COD"]))["UNIDADE_SIGLA"].split()[0]   # '511o'
    partes = cabeca.split()
    opcoes = [cabeca, " ".join(partes[:2]), cabeca + "/" + om_num.rstrip("oa"), sigla,
              cabeca.replace("1a", "1ª").replace("2a", "2ª").replace("3a", "3ª"),
              cabeca.replace("1a ", "1 ").replace("2a ", "2 ").replace("3a ", "3 ")]
    return rng.choice(opcoes)


def _estado_inicial(ctx: Contexto, u: dict) -> dict:
    rng = ctx.rng
    esc = u["ESCALAO_COD"]
    sigla = u["UNIDADE_SIGLA"]
    if esc == "SU":
        ef = rng.randint(115, 140) if "Cia" in sigla else rng.randint(85, 110)
        vtr = rng.randint(10, 16)
    else:
        ef = rng.randint(26, 38)
        vtr = rng.randint(3, 6)
    return {"ef_prev": ef, "ausentes": rng.randint(0, 6), "bx_cmb": 0, "bx_ncmb": 0, "evac": 0,
            "vtr_total": vtr, "vtr_op": vtr - rng.randint(0, 2), "comb": rng.uniform(70, 100),
            "sup": {k: rng.uniform(3, 5) for k in ("I", "III", "V", "VIII")}}


def _avancar(ctx: Contexto, e: dict, turno_idx: int):
    """Uma remessa: consumo, baixas e reabastecimento."""
    rng = ctx.rng
    e["comb"] = max(10, e["comb"] - rng.uniform(4, 14))
    for k in e["sup"]:
        e["sup"][k] = max(0.3, e["sup"][k] - rng.uniform(0.2, 0.5))
    if rng.random() < 0.18:                       # reabastecimento
        e["comb"] = min(100, e["comb"] + rng.uniform(25, 60))
        for k in e["sup"]:
            e["sup"][k] = min(6, e["sup"][k] + rng.uniform(1, 3))
    if rng.random() < 0.06:
        e["bx_cmb"] += rng.randint(1, 3)          # exercicio: baixas simuladas
    if rng.random() < 0.10:
        e["bx_ncmb"] += 1
    if rng.random() < 0.12:
        e["evac"] += 1
    if rng.random() < 0.15 and e["evac"] > 0:
        e["evac"] -= 1
    if rng.random() < 0.10:
        e["vtr_op"] = max(0, e["vtr_op"] - 1)
    elif rng.random() < 0.15:
        e["vtr_op"] = min(e["vtr_total"], e["vtr_op"] + 1)
    e["ausentes"] = max(0, e["ausentes"] + rng.randint(-1, 1))


def _valores(ctx: Contexto, e: dict, versao: int, nec: str) -> list:
    v = [e["ef_prev"], e["ef_prev"] - e["ausentes"] - e["bx_cmb"] - e["bx_ncmb"] - e["evac"],
         e["bx_cmb"], e["bx_ncmb"], e["evac"], e["vtr_op"], e["vtr_total"], round(e["comb"], 1),
         round(e["sup"]["I"], 1), round(e["sup"]["III"], 1), round(e["sup"]["V"], 1), round(e["sup"]["VIII"], 1), nec]
    if versao == 2:
        v.append(ctx.rng.randint(200, 2500))
    return v


def _xlsx(caminho, om_sigla, data_txt, turno, cabecalho, linhas):
    wb = Workbook()
    ws = wb.active
    ws.title = "RELPER"
    ws["A1"] = "RELATÓRIO PERIÓDICO DE PESSOAL E MATERIAL"
    ws["A2"] = "Operação Perseu 2024 — 51ª Bda Inf Mtz"
    ws["A3"] = f"OM: {om_sigla}"
    ws["A4"] = f"Data: {data_txt}    Turno: {turno}"
    ws.append([])
    ws.append(cabecalho)
    for l in linhas:
        ws.append(l)
    ws.column_dimensions["A"].width = 28
    wb.save(caminho)


def _pdf(ctx: Contexto, caminho, om_sigla, data_txt, turno, cabecalho, linhas):
    img, d = pagina()
    texto_centrado(d, 110, "RELATÓRIO PERIÓDICO DE PESSOAL E MATERIAL", fonte(40, negrito=True))
    texto_centrado(d, 165, "Operação Perseu 2024 — 51ª Bda Inf Mtz", fonte(26))
    d.text((80, 250), f"OM: {om_sigla}", font=fonte(26), fill="black")
    d.text((80, 295), f"Data: {data_txt}      Turno: {turno}", font=fonte(26), fill="black")
    n = len(cabecalho)
    larguras = [250] + [(1574 - 250) // (n - 1)] * (n - 1)
    y = desenhar_tabela(d, 40, 370, larguras, cabecalho, linhas, fonte(14, negrito=True), fonte(18), alt_cab=64, alt_lin=54, margem=5)
    d.text((80, y + 120), "Confere.", font=fonte(24), fill="black")
    assinar(d, 80, y + 200, ctx.rng)
    d.text((80, y + 250), f"Cmt {om_sigla}", font=fonte(22), fill="black")
    salvar_pdf([escanear(img, ctx.rng, ctx.np_rng)], caminho)


def gerar(ctx: Contexto) -> dict:
    rng = ctx.rng
    pasta = ctx.pasta("relper")
    pasta_errada = ctx.landing.parent / "perseu_2023" / "relper"
    oms = ctx.por_escalao("OM")
    estados = {u["UNIDADE_COD"]: _estado_inicial(ctx, u) for om in oms for u in _linhas_da_om(ctx, om["UNIDADE_COD"])}
    arquivos, linhas_total, ausentes = 0, 0, 0
    defeito_pasta_feito = False
    for dia in ctx.operacao.dias:
        for turno, hora in TURNOS:
            for om in oms:
                om_cod = om["UNIDADE_COD"]
                unidades = _linhas_da_om(ctx, om_cod)
                for u in unidades:
                    _avancar(ctx, estados[u["UNIDADE_COD"]], 0)
                if rng.random() < 0.08:
                    ausentes += 1
                    ctx.registrar_defeito("RELPER", f"{om_cod} {dia} {turno}", "turno_ausente",
                                          "a OM nao remeteu o relatorio deste turno; a ausencia e dado")
                    continue
                versao = 2 if om_cod in OM_V2 else 1
                cabecalho = list(CABECALHO_V2 if versao == 2 else CABECALHO_V1)
                extra = OM_COLUNA_EXTRA.get(om_cod)
                if extra:
                    cabecalho.append(extra)
                linhas = []
                for u in unidades:
                    nec = rng.choice(NECESSIDADES)
                    if nec == "Fz" and rng.random() < 0.35:
                        nec = "Fuz"
                        ctx.registrar_defeito("RELPER", f"{om_cod} {dia} {turno} {u['UNIDADE_COD']}", "abreviatura_trocada",
                                              "'Fuz' (Fuzileiro) escrito no lugar de 'Fz' (Fuzil): material vira pessoal")
                    linha = [_grafia_fracao(ctx, u)] + _valores(ctx, estados[u["UNIDADE_COD"]], versao, nec)
                    if extra:
                        linha.append(rng.choice(OBS))
                    linhas.append(linha)
                t = local(dia, hora, rng.randint(0, 59)) + timedelta(minutes=rng.randint(-30, 40))
                data_txt = dia.strftime("%d/%m/%Y")
                ext = "pdf" if om_cod in OM_PDF else "xlsx"
                destino = pasta
                nome = f"relper_{om_cod.lower()}_{dia.isoformat()}_{turno.lower()}.{ext}"
                if ext == "xlsx" and not defeito_pasta_feito and dia == ctx.operacao.dias[3]:
                    destino = pasta_errada
                    destino.mkdir(parents=True, exist_ok=True)
                    defeito_pasta_feito = True
                    ctx.registrar_defeito("RELPER", (destino / nome).relative_to(ctx.saida), "pasta_errada",
                                          "arquivo pousou em landing/perseu_2023/; o sidecar diz PERSEU_2024")
                caminho = destino / nome
                sidecar = {"operacao": ctx.operacao.cod, "om_remetente": om["UNIDADE_SIGLA"],
                           "data_hora": iso(t), "turno": turno}
                if ext == "xlsx":
                    _xlsx(caminho, om["UNIDADE_SIGLA"], data_txt, turno, cabecalho, linhas)
                else:
                    _pdf(ctx, caminho, om["UNIDADE_SIGLA"], data_txt, turno, cabecalho, linhas)
                escrever_sidecar(caminho, sidecar)
                arquivos += 1

                # Algumas remessas vao NOS DOIS formatos: a OM manda a planilha e,
                # para o arquivo, tambem o formulario impresso, assinado e
                # digitalizado. Os mesmos numeros por dois caminhos — leitura
                # direta de um lado, OCR do outro — e por isso o lado da planilha
                # e a VERDADE por construcao, sem precisar de gabarito a parte.
                if ext == "xlsx" and rng.random() < 0.12:
                    gemeo = caminho.with_name(caminho.stem + "_digitalizado.pdf")
                    _pdf(ctx, gemeo, om["UNIDADE_SIGLA"], data_txt, turno, cabecalho, linhas)
                    escrever_sidecar(gemeo, sidecar)
                    ctx.registrar_gabarito("RELPER", gemeo.relative_to(ctx.saida),
                                           "planilha_equivalente", caminho.relative_to(ctx.saida))
                    arquivos += 1
                linhas_total += len(linhas)
    r = {"arquivos_binarios": arquivos, "linhas": linhas_total, "turnos_ausentes": ausentes}
    ctx.resumo["RELPER"] = r
    return r


if __name__ == "__main__":
    from gerar_tudo import executar_um
    executar_um("RELPER", gerar)
